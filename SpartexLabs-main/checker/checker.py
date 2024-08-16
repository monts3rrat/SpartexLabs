import blocksmith
from openpyxl import load_workbook
import asyncio


async def generate_address(path_to_xlsx):
    workbook = load_workbook(path_to_xlsx)
    sheet = workbook.active

    addresses = []

    i = 0
    while True:
        i += 1
        private_key = sheet.cell(row=i + 1, column=2).value
        if (
                sheet.cell(row=i + 1, column=2).value == ""
                or sheet.cell(row=i + 1, column=2).value is None
        ):
            break
        addresses.append(blocksmith.EthereumWallet.generate_address(private_key))

    return addresses


async def generate_address_zora(path_to_xlsx):
    workbook = load_workbook(path_to_xlsx)
    sheet = workbook.active

    addresses = []

    i = 0
    while True:
        i += 1
        private_key = sheet.cell(row=i + 1, column=1).value
        if (
                sheet.cell(row=i + 1, column=1).value == ""
                or sheet.cell(row=i + 1, column=1).value is None
        ):
            break
        addresses.append(blocksmith.EthereumWallet.generate_address(private_key))

    return addresses
